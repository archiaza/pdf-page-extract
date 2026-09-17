"""Tests for pagextract. Run with:  python -m pytest -v"""

from __future__ import annotations

import io
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from pagextract import (  # noqa: E402
    Block,
    Document,
    ExtractionError,
    UnsupportedFormat,
    extract,
    extract_docx,
    extract_pdf,
)

# --------------------------------------------------------------------------- #
# fixtures
# --------------------------------------------------------------------------- #

reportlab = pytest.importorskip("reportlab", reason="reportlab needed to build sample PDFs")
docx_mod = pytest.importorskip("docx", reason="python-docx needed to build sample DOCX")


def make_pdf(path: Path, pages: int) -> Path:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path), pagesize=A4)
    for i in range(1, pages + 1):
        c.drawString(72, 800, f"PAGE MARKER {i}")
        c.drawString(72, 780, "Sample document body. No confidential data.")
        c.showPage()
    c.save()
    return path


def make_docx(path: Path) -> Path:
    from docx import Document as D

    d = D()
    d.add_paragraph("EXPLANATORY NOTE")
    d.add_paragraph("Room height is 2.5 m.")
    d.add_paragraph("")  # deliberately blank
    t = d.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text = "Room"
    t.rows[0].cells[1].text = "Area, m2"
    t.rows[1].cells[0].text = "Office"
    t.rows[1].cells[1].text = "18.4"
    d.save(str(path))
    return path


@pytest.fixture
def pdf5(tmp_path: Path) -> Path:
    return make_pdf(tmp_path / "five.pdf", 5)


@pytest.fixture
def pdf12(tmp_path: Path) -> Path:
    return make_pdf(tmp_path / "twelve.pdf", 12)


@pytest.fixture
def docx1(tmp_path: Path) -> Path:
    return make_docx(tmp_path / "note.docx")


# --------------------------------------------------------------------------- #
# Block / Document
# --------------------------------------------------------------------------- #

def test_block_is_empty():
    assert Block("", "p. 1", "page", 1).is_empty
    assert Block("   \n ", "p. 1", "page", 1).is_empty
    assert not Block("text", "p. 1", "page", 1).is_empty


def test_block_len():
    assert len(Block("abcd", "p. 1", "page", 1)) == 4


def test_document_helpers():
    d = Document(source="x", fmt="pdf", blocks=[
        Block("alpha", "p. 1", "page", 1),
        Block("", "p. 2", "page", 2),
        Block("beta ALPHA", "p. 3", "page", 3),
    ])
    assert len(d) == 3
    assert len(d.non_empty()) == 2
    assert d.char_count == len("alpha") + len("beta ALPHA")
    assert d.text == "alpha\n\nbeta ALPHA"
    assert d.locate("alpha") == ["p. 1", "p. 3"]           # case-insensitive
    assert d.locate("alpha", case_sensitive=True) == ["p. 1"]
    assert [b.locator for b in d] == ["p. 1", "p. 2", "p. 3"]


def test_document_to_dict_roundtrip():
    d = Document(source="x", fmt="pdf", blocks=[Block("a", "p. 1", "page", 1)])
    out = d.to_dict()
    assert out["format"] == "pdf"
    assert out["blocks"][0]["locator"] == "p. 1"
    assert out["char_count"] == 1


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #

def test_pdf_every_page_is_a_block(pdf5):
    doc = extract_pdf(pdf5)
    assert doc.fmt == "pdf"
    assert len(doc) == 5
    assert doc.total_units == 5
    assert doc.truncated is False


def test_pdf_locators_are_sequential_and_one_based(pdf5):
    doc = extract_pdf(pdf5)
    assert [b.locator for b in doc] == [f"p. {i}" for i in range(1, 6)]
    assert [b.index for b in doc] == [1, 2, 3, 4, 5]


def test_pdf_text_lands_on_the_right_page(pdf5):
    doc = extract_pdf(pdf5)
    for i in range(1, 6):
        assert doc.locate(f"PAGE MARKER {i}") == [f"p. {i}"]


def test_pdf_max_pages_truncates_and_says_so(pdf12):
    doc = extract_pdf(pdf12, max_pages=4)
    assert len(doc) == 4
    assert doc.truncated is True
    assert doc.total_units == 12          # real size stays visible
    assert any("stopped at page 4 of 12" in w for w in doc.warnings)


def test_pdf_no_limit_reads_everything(pdf12):
    doc = extract_pdf(pdf12)
    assert len(doc) == 12
    assert doc.truncated is False
    assert doc.locate("PAGE MARKER 12") == ["p. 12"]


def test_pdf_max_pages_above_total_is_not_truncation(pdf5):
    doc = extract_pdf(pdf5, max_pages=999)
    assert len(doc) == 5
    assert doc.truncated is False


def test_pdf_max_pages_must_be_positive(pdf5):
    with pytest.raises(ValueError):
        extract_pdf(pdf5, max_pages=0)


def test_pdf_accepts_a_stream(pdf5):
    stream = io.BytesIO(pdf5.read_bytes())
    stream.name = "five.pdf"
    doc = extract_pdf(stream)
    assert len(doc) == 5


def test_pdf_broken_file_raises_not_silently_empty(tmp_path):
    bad = tmp_path / "broken.pdf"
    bad.write_bytes(b"%PDF-1.4 this is not a real pdf")
    with pytest.raises(ExtractionError):
        extract_pdf(bad)


def test_missing_file_raises(tmp_path):
    with pytest.raises(FileNotFoundError):
        extract_pdf(tmp_path / "nope.pdf")


# --------------------------------------------------------------------------- #
# DOCX
# --------------------------------------------------------------------------- #

def test_docx_reads_paragraphs(docx1):
    doc = extract_docx(docx1)
    assert doc.fmt == "docx"
    assert doc.locate("EXPLANATORY NOTE") == ["para 1"]
    assert doc.locate("2.5 m") == ["para 2"]


def test_docx_reads_tables_by_default(docx1):
    doc = extract_docx(docx1)
    tables = [b for b in doc if b.kind == "table"]
    assert len(tables) == 1
    assert "Office" in tables[0].text
    assert "18.4" in tables[0].text
    assert tables[0].locator == "table 1"


def test_docx_tables_can_be_skipped(docx1):
    doc = extract_docx(docx1, include_tables=False)
    assert [b for b in doc if b.kind == "table"] == []
    assert doc.locate("Office") == []


def test_docx_keeps_blank_paragraph_but_marks_it_empty(docx1):
    doc = extract_docx(docx1)
    blanks = [b for b in doc if b.kind == "paragraph" and b.is_empty]
    assert blanks, "the blank paragraph should still be present"


def test_docx_warns_that_locators_are_not_pages(docx1):
    doc = extract_docx(docx1)
    assert any("not page numbers" in w for w in doc.warnings)


def test_docx_broken_file_raises(tmp_path):
    bad = tmp_path / "broken.docx"
    bad.write_bytes(b"PK\x03\x04 not really a docx")
    with pytest.raises(ExtractionError):
        extract_docx(bad)


# --------------------------------------------------------------------------- #
# dispatch
# --------------------------------------------------------------------------- #

def test_extract_dispatches_by_extension(pdf5, docx1):
    assert extract(pdf5).fmt == "pdf"
    assert extract(docx1).fmt == "docx"


def test_extract_forwards_kwargs(pdf12, docx1):
    assert len(extract(pdf12, max_pages=3)) == 3
    assert [b for b in extract(docx1, include_tables=False) if b.kind == "table"] == []


def test_legacy_doc_is_rejected_with_advice(tmp_path):
    f = tmp_path / "old.doc"
    f.write_bytes(b"\xd0\xcf\x11\xe0")
    with pytest.raises(UnsupportedFormat, match="convert"):
        extract(f)


def test_unknown_extension_is_rejected(tmp_path):
    f = tmp_path / "data.txt"
    f.write_text("hello")
    with pytest.raises(UnsupportedFormat):
        extract(f)


# --------------------------------------------------------------------------- #
# CLI
# --------------------------------------------------------------------------- #

def test_cli_summary(pdf5, capsys):
    from pagextract.cli import main
    assert main([str(pdf5), "--format", "summary", "--quiet"]) == 0
    out = capsys.readouterr().out
    assert "blocks      : 5" in out


def test_cli_find_reports_locator(pdf5, capsys):
    from pagextract.cli import main
    assert main([str(pdf5), "--find", "PAGE MARKER 3", "--quiet"]) == 0
    assert "--- p. 3 ---" in capsys.readouterr().out


def test_cli_json_is_valid(pdf5, capsys):
    import json
    from pagextract.cli import main
    assert main([str(pdf5), "--format", "json", "--quiet"]) == 0
    payload = json.loads(capsys.readouterr().out)
    assert payload["format"] == "pdf"
    assert len(payload["blocks"]) == 5


def test_cli_missing_file_returns_1(tmp_path, capsys):
    from pagextract.cli import main
    assert main([str(tmp_path / "nope.pdf"), "--quiet"]) == 1
    assert "error:" in capsys.readouterr().err
