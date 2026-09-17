"""
Working with .docx, including tables.

Run:  python examples/find_in_docx.py

Shows two things that trip people up:
  1. tables carry real content and are easy to lose;
  2. .docx has no page numbers until it is rendered.
"""

import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from pagextract import extract  # noqa: E402


def build_sample_docx(path: Path) -> Path:
    from docx import Document as D

    d = D()
    d.add_paragraph("EXPLANATORY NOTE")
    d.add_paragraph("Section 1. General data.")
    d.add_paragraph("The building has three storeys above ground.")

    table = d.add_table(rows=4, cols=3)
    rows = [
        ("No.", "Room", "Area, m2"),
        ("1", "Office", "18.4"),
        ("2", "Meeting room", "24.0"),
        ("3", "Server room", "9.6"),
    ]
    for r, values in enumerate(rows):
        for c, value in enumerate(values):
            table.rows[r].cells[c].text = value

    d.add_paragraph("Section 2. Structural solutions.")
    d.save(str(path))
    return path


def main() -> None:
    tmp = Path(tempfile.mkdtemp())
    docx = build_sample_docx(tmp / "note.docx")

    doc = extract(docx)
    print(f"blocks: {len(doc)}  ({len(doc.non_empty())} non-empty)\n")

    by_kind: dict[str, int] = {}
    for b in doc:
        by_kind[b.kind] = by_kind.get(b.kind, 0) + 1
    print(f"by kind: {by_kind}\n")

    # Table content is preserved -- this is where specifications usually live.
    print("Where is 'Server room'?")
    print(f"  {doc.locate('Server room')}\n")

    tables = [b for b in doc if b.kind == "table"]
    if tables:
        print("Table 1 content:")
        for line in tables[0].text.splitlines():
            print(f"  {line}")
    print()

    # Compare: dropping tables loses that content entirely.
    without = extract(docx, include_tables=False)
    print(f"include_tables=False -> 'Server room' found at: "
          f"{without.locate('Server room') or 'nowhere (content lost)'}")
    print()

    for w in doc.warnings:
        print(f"note: {w}")


if __name__ == "__main__":
    main()
